# 05_supp_table_markers.py
# Supplementary Table S3: the marker genes, their source (Greening lab databases or KEGG KOfam), inclusion threshold,
# and the number of proteins and MAGs counted among the 513 dereplicated MAGs.
# Run from the repository root after make_supp_tables.R (see README.md).
import csv, collections, openpyxl
from openpyxl.styles import Font, Alignment

mags = {r["MAG"] for r in csv.DictReader(open("results/supplementary/TableS2_MAGs.tsv"), delimiter="\t")}
assert len(mags) == 513
prot = collections.Counter(); hasm = collections.defaultdict(set)
for r in csv.DictReader(open("data/functional/protein_marker_calls.tsv"), delimiter="\t"):
    mag = r["protein"].split("~")[0]
    if mag in mags:
        m = {"btuC": "BtuCDF", "btuD": "BtuCDF", "btuF": "BtuCDF"}.get(r["marker"], r["marker"])
        prot[(m, r["marker"])] += 1; hasm[m].add(mag)
mat = {r[0]: r[1:] for r in csv.reader(open("data/functional/gene_bin_mat.csv"))}
hdr = mat.pop("gene")
complete = {m: sum(1 for mag, v in zip(hdr, vals) if mag in mags and int(v) > 0) for m, vals in mat.items()}

SYM = {"PEPC": "ppc", "AcnB": "acnB", "CA": "carbonic anhydrase", "MDH": "mxaF/xoxF", "PsbA": "psbA", "PufM": "pufM",
       "BluB": "bluB", "BtuCDF": "btuCDF", "CobA": "cobA", "CobC": "cobC", "CobU": "cobU", "CbiG": "cbiG", "ArsC": "arsC",
       "RHO": "rhodopsin", "NiFe": "[NiFe]-hydrogenase", "FeFe": "[FeFe]-hydrogenase", "FCC": "fccB"}
sym = lambda m: SYM.get(m, m[0].lower() + m[1:] if m[0].isupper() and m[1:2].islower() else m)
thr = {r["knum"]: r for r in csv.DictReader(open("database/kofam_thresholds.tsv"), delimiter="\t")}

rows = []
for r in csv.DictReader(open("database/greening_markers.tsv"), delimiter="\t"):
    m = r["marker"]
    rows.append([sym(m), r["pathway"], "Greening lab databases, DIAMOND blastp", r["greening_set"],
                 f"identity >= {r['min_identity_pct']}%, query coverage >= 80%", int(r["reference_sequences"]),
                 prot[(m, m)], len(hasm[m])])
# one row per marker; markers built from several KOs (carbonic anhydrases, mxaF/xoxF, arsC, btuCDF) list every KO
kos = collections.defaultdict(list)
for r in csv.DictReader(open("database/kofam_markers.tsv"), delimiter="\t"):
    if r["marker"] != "SoxB_check":
        m = {"btuC": "BtuCDF", "btuD": "BtuCDF", "btuF": "BtuCDF"}.get(r["marker"], r["marker"])
        kos[(m, r["pathway"])].append(r)
for (m, pw), rs in kos.items():
    ref = "; ".join(f"{r['ko']} ({thr[r['ko']]['definition']})" for r in rs)
    cut = "; ".join(f"{r['ko']}: {thr[r['ko']]['score_type']} score >= {thr[r['ko']]['threshold']}" for r in rs)
    n_prot = sum(prot[(m, r["marker"])] for r in {r["marker"]: r for r in rs}.values())
    rows.append([sym(m), pw, "KEGG KOfam, hmmsearch", ref, cut, None, n_prot,
                 complete.get("BtuCDF", 0) if m == "BtuCDF" else len(hasm[m])])
order = ["Respiration", "Carbon core", "Cobalamin (B12)", "DMSP and methylated compounds", "Nitrogen cycle",
         "Sulfur cycle", "Methane cycling", "Phototrophy", "Hydrogenases", "Detoxification and other",
         "Not used in pathway analyses"]
rows.sort(key=lambda x: (order.index(x[1]), x[0].lower()))

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Supplementary Table S3"
ws.append(["Supplementary Table S3. Marker genes used to estimate functional potential. Proteins predicted from the MAGs "
           "were searched against the Greening lab metabolic marker gene databases (DIAMOND blastp, best hit among all 51 "
           "marker sets, Greening lab identity thresholds for MAGs) or KEGG KOfam profile HMMs (hmmsearch, KEGG adaptive "
           "score thresholds; full-sequence or best-domain score as specified by KEGG). Counts refer to the 513 dereplicated "
           "MAGs; for btuCDF, the MAG count is the number of MAGs with a complete transporter (btuC, btuD and btuF). "
           "Not included: isiB (iron-stress flavodoxin, not a cobalamin gene), dmoA (no KEGG threshold for K16967), and "
           "amoB and amoC (KEGG K10945 and K10946 do not separate ammonia from methane monooxygenase subunits)."])
ws.append([])
ws.append(["Gene", "Pathway in this study", "Method", "Reference set or KEGG KO", "Inclusion threshold",
           "Reference sequences", "Proteins counted", "MAGs with the gene"])
for c in ws[3]:
    c.font = Font(bold=True)
for r in rows:
    ws.append(r)
for col, wd in zip("ABCDEFGH", (20, 30, 36, 62, 34, 12, 12, 12)):
    ws.column_dimensions[col].width = wd
ws.merge_cells("A1:H1"); ws["A1"].alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[1].height = 92
wb.save("results/supplementary/Supplementary_Table_S3_markers.xlsx")
print("rows:", len(rows), "| Greening:", sum(r[2].startswith("Greening") for r in rows), "| KOfam:", sum(r[2].startswith("KEGG") for r in rows))
